from __future__ import annotations

import json
import os
import subprocess
from concurrent.futures import Future, ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from time import perf_counter

import numpy as np
import psutil
from PyQt6.QtCore import QObject, QTimer

try:
    import torch
except ImportError:
    torch = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
except ImportError:
    Workbook = None
    Alignment = None
    Font = None


class PerformanceSessionLogger(QObject):
    """Collect performance metrics for single- and multi-person workouts.

    Each mode is measured as an independent logger session. When a mode ends,
    its result is checkpointed immediately. The XLSX report is then regenerated
    from all completed mode results retained for the same test run. Therefore:

    - after the single-person workout, the report already contains single data;
    - after the multi-person workout, the same report contains both modes;
    - data from the single-person workout is not discarded;
    - workout-guide time between exercises is excluded through pause/resume.
    """

    def __init__(
        self,
        output_dir: Path,
        parent=None,
        resource_interval_ms: int = 1000,
    ):
        super().__init__(parent)

        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.process = psutil.Process(os.getpid())
        self.logical_cpu_count = psutil.cpu_count(logical=True) or 1

        self.resource_timer = QTimer(self)
        self.resource_timer.setInterval(resource_interval_ms)
        self.resource_timer.timeout.connect(self._sample_resources)

        # nvidia-smi can take tens or hundreds of milliseconds. Run it outside
        # the GUI thread so resource monitoring does not distort FPS/latency.
        self._gpu_executor = ThreadPoolExecutor(
            max_workers=1,
            thread_name_prefix="movion-perf-gpu",
        )
        self._gpu_future: Future | None = None
        self._last_gpu_percent: float | None = None
        self._last_gpu_memory_mb: float | None = None

        # Data retained across the single and multi sessions in one app run.
        self._run_id: str | None = None
        self._completed_results: dict[str, dict] = {}
        self.output_path: Path | None = None
        self.fallback_json_path: Path | None = None
        self.state_path = self.output_dir / ".performance_test_state.json"

        self._reset_current_session()

    def _reset_current_session(self):
        self.running = False
        self.segment_active = False

        self.mode: str | None = None
        self.metadata: dict = {}

        self.started_datetime: datetime | None = None
        self.segment_started_at: float | None = None
        self.active_duration_seconds = 0.0

        self.input_frame_count = 0
        self.processed_frame_count = 0

        self.model_inference_count = 0
        self.model_inference_per_user = {
            "user1": 0,
            "user2": 0,
        }

        self.latency_samples_ms: list[float] = []
        self.resource_samples: list[dict] = []

        # Do not carry a GPU sample from the previous workout mode.
        self._gpu_future = None
        self._last_gpu_percent = None
        self._last_gpu_memory_mb = None

    def _restore_incomplete_run(self) -> bool:
        if not self.state_path.exists():
            return False

        try:
            state = json.loads(
                self.state_path.read_text(encoding="utf-8")
            )
            run_id = state.get("run_id")
            completed = state.get("sessions") or {}

            if not run_id or "single" not in completed or "multi" in completed:
                return False

            self._run_id = str(run_id)
            self._completed_results = dict(completed)
            self.output_path = (
                self.output_dir
                / f"hasil_pengujian_performa_{self._run_id}.xlsx"
            )
            self.fallback_json_path = (
                self.output_dir
                / f"hasil_pengujian_performa_{self._run_id}.json"
            )
            return True
        except (OSError, ValueError, TypeError):
            return False

    def _save_state(self):
        if self._run_id is None:
            return

        state = {
            "run_id": self._run_id,
            "sessions": self._completed_results,
        }
        self.state_path.write_text(
            json.dumps(state, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )

    def _begin_report_run_if_needed(self, mode: str):
        if mode == "multi" and self._restore_incomplete_run():
            return

        new_run = (
            self._run_id is None
            or mode == "single"
            or {"single", "multi"}.issubset(self._completed_results)
        )

        if not new_run:
            return

        self._run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        self._completed_results = {}
        self.output_path = (
            self.output_dir
            / f"hasil_pengujian_performa_{self._run_id}.xlsx"
        )
        self.fallback_json_path = (
            self.output_dir
            / f"hasil_pengujian_performa_{self._run_id}.json"
        )
        self._save_state()

    def start(self, mode: str, metadata: dict | None = None):
        if self.running:
            return

        normalized_mode = "multi" if mode == "multi" else "single"
        self._begin_report_run_if_needed(normalized_mode)
        self._reset_current_session()

        self.running = True
        self.mode = normalized_mode
        self.metadata = dict(metadata or {})
        self.started_datetime = datetime.now()

        self.process.cpu_percent(interval=None)

        if torch is not None and torch.cuda.is_available():
            try:
                torch.cuda.reset_peak_memory_stats()
            except Exception:
                pass

        self.resource_timer.start()

        print(
            "[PERF_TEST_STARTED]",
            f"mode={self.mode}",
            f"time={self.started_datetime.isoformat(timespec='seconds')}",
        )

    def resume(self):
        if not self.running or self.segment_active:
            return

        self.segment_active = True
        self.segment_started_at = perf_counter()
        print("[PERF_SEGMENT_RESUMED]")

    def pause(self):
        if not self.running or not self.segment_active:
            return

        now = perf_counter()
        if self.segment_started_at is not None:
            self.active_duration_seconds += now - self.segment_started_at

        self.segment_active = False
        self.segment_started_at = None
        print("[PERF_SEGMENT_PAUSED]")

    def _active_duration(self) -> float:
        duration = self.active_duration_seconds
        if self.segment_active and self.segment_started_at is not None:
            duration += perf_counter() - self.segment_started_at
        return max(0.0, duration)

    def record_input_frame(
        self,
        width: int | None = None,
        height: int | None = None,
    ):
        if not self.running or not self.segment_active:
            return

        self.input_frame_count += 1

        if width and height and not self.metadata.get("camera_resolution"):
            self.metadata["camera_resolution"] = f"{int(width)}x{int(height)}"

    def record_processed_frame(self, end_to_end_latency_ms: float):
        if not self.running or not self.segment_active:
            return

        try:
            latency = float(end_to_end_latency_ms)
        except (TypeError, ValueError):
            return

        if not np.isfinite(latency) or latency < 0:
            return

        self.processed_frame_count += 1
        self.latency_samples_ms.append(latency)

    def record_model_inference(self, user_key: str):
        if not self.running or not self.segment_active:
            return

        self.model_inference_count += 1
        if user_key in self.model_inference_per_user:
            self.model_inference_per_user[user_key] += 1

    @staticmethod
    def _query_nvidia_smi() -> tuple[float | None, float | None]:
        """Return device GPU utilization (%) and used memory (MB).

        This uses the NVIDIA driver utility directly and therefore does not
        require the optional nvidia-ml-py/pynvml package. Values are device-wide,
        so other GPU-heavy programs should be closed during the experiment.
        """
        command = [
            "nvidia-smi",
            "--query-gpu=utilization.gpu,memory.used",
            "--format=csv,noheader,nounits",
            "--id=0",
        ]

        kwargs = {
            "capture_output": True,
            "text": True,
            "timeout": 1.5,
            "check": False,
        }
        if os.name == "nt":
            kwargs["creationflags"] = getattr(
                subprocess,
                "CREATE_NO_WINDOW",
                0,
            )

        try:
            result = subprocess.run(command, **kwargs)
            if result.returncode != 0 or not result.stdout.strip():
                return None, None

            first_line = result.stdout.strip().splitlines()[0]
            fields = [field.strip() for field in first_line.split(",")]
            if len(fields) < 2:
                return None, None

            return float(fields[0]), float(fields[1])
        except (OSError, ValueError, subprocess.SubprocessError):
            return None, None

    def _sample_resources(self):
        if not self.running or not self.segment_active:
            return

        try:
            raw_process_cpu = self.process.cpu_percent(interval=None)
            cpu_percent = raw_process_cpu / self.logical_cpu_count
            ram_mb = self.process.memory_info().rss / (1024.0 ** 2)
        except Exception as exc:
            print("[PERF_RESOURCE_CPU_RAM_WARN]", exc)
            return

        # Consume a completed GPU query without blocking this QTimer callback.
        if self._gpu_future is not None and self._gpu_future.done():
            try:
                (
                    self._last_gpu_percent,
                    self._last_gpu_memory_mb,
                ) = self._gpu_future.result()
            except Exception:
                self._last_gpu_percent = None
                self._last_gpu_memory_mb = None
            finally:
                self._gpu_future = None

        # Start the next device query in the background. The latest completed
        # values are used for the current sample.
        if self._gpu_future is None:
            self._gpu_future = self._gpu_executor.submit(
                self._query_nvidia_smi
            )

        gpu_percent = self._last_gpu_percent
        gpu_memory_mb = self._last_gpu_memory_mb

        torch_gpu_memory_mb = None
        if torch is not None and torch.cuda.is_available():
            try:
                torch_gpu_memory_mb = (
                    torch.cuda.memory_reserved() / (1024.0 ** 2)
                )
            except Exception:
                torch_gpu_memory_mb = None

        self.resource_samples.append(
            {
                "active_time_s": self._active_duration(),
                "cpu_percent": cpu_percent,
                "ram_mb": ram_mb,
                "gpu_percent": gpu_percent,
                "gpu_memory_mb": gpu_memory_mb,
                "torch_gpu_memory_mb": torch_gpu_memory_mb,
            }
        )

    @staticmethod
    def _average(values):
        valid = [
            float(value)
            for value in values
            if value is not None and np.isfinite(value)
        ]
        return float(np.mean(valid)) if valid else None

    @staticmethod
    def _maximum(values):
        valid = [
            float(value)
            for value in values
            if value is not None and np.isfinite(value)
        ]
        return float(np.max(valid)) if valid else None

    def _build_summary(self) -> dict:
        duration = self._active_duration()
        latencies = np.asarray(self.latency_samples_ms, dtype=np.float64)

        if latencies.size:
            latency_summary = {
                "latency_mean_ms": float(np.mean(latencies)),
                "latency_median_ms": float(np.median(latencies)),
                "latency_p95_ms": float(np.percentile(latencies, 95)),
                "latency_min_ms": float(np.min(latencies)),
                "latency_max_ms": float(np.max(latencies)),
            }
        else:
            latency_summary = {
                "latency_mean_ms": None,
                "latency_median_ms": None,
                "latency_p95_ms": None,
                "latency_min_ms": None,
                "latency_max_ms": None,
            }

        input_fps = (
            self.input_frame_count / duration if duration > 0 else 0.0
        )
        effective_fps = (
            self.processed_frame_count / duration if duration > 0 else 0.0
        )
        skipping_ratio = (
            (1.0 - self.processed_frame_count / self.input_frame_count) * 100.0
            if self.input_frame_count > 0
            else 0.0
        )
        inference_hz = (
            self.model_inference_count / duration if duration > 0 else 0.0
        )

        cpu_values = [sample["cpu_percent"] for sample in self.resource_samples]
        ram_values = [sample["ram_mb"] for sample in self.resource_samples]
        gpu_values = [sample["gpu_percent"] for sample in self.resource_samples]
        gpu_memory_values = [
            sample["gpu_memory_mb"] for sample in self.resource_samples
        ]
        torch_memory_values = [
            sample["torch_gpu_memory_mb"] for sample in self.resource_samples
        ]

        summary = {
            "mode": self.mode,
            "started_at": (
                self.started_datetime.isoformat(timespec="seconds")
                if self.started_datetime
                else None
            ),
            "duration_s": duration,
            "input_frames": self.input_frame_count,
            "processed_frames": self.processed_frame_count,
            "input_fps": input_fps,
            "effective_fps": effective_fps,
            "frame_skipping_ratio_percent": skipping_ratio,
            "model_inference_count": self.model_inference_count,
            "model_inference_hz": inference_hz,
            "user1_inference_count": self.model_inference_per_user["user1"],
            "user2_inference_count": self.model_inference_per_user["user2"],
            "cpu_mean_percent": self._average(cpu_values),
            "cpu_max_percent": self._maximum(cpu_values),
            "gpu_mean_percent": self._average(gpu_values),
            "gpu_max_percent": self._maximum(gpu_values),
            "ram_mean_mb": self._average(ram_values),
            "ram_max_mb": self._maximum(ram_values),
            "gpu_memory_max_mb": self._maximum(gpu_memory_values),
            "torch_gpu_memory_max_mb": self._maximum(torch_memory_values),
        }
        summary.update(latency_summary)
        return summary

    @staticmethod
    def _round_or_none(value, digits: int = 3):
        if isinstance(value, float):
            return round(value, digits)
        return value

    def _write_xlsx(self):
        if Workbook is None or self.output_path is None:
            return False

        workbook = Workbook()
        summary_ws = workbook.active
        summary_ws.title = "Ringkasan"

        summary_ws["A1"] = "HASIL PENGUJIAN PERFORMA APLIKASI"
        summary_ws["A1"].font = Font(bold=True, size=14)
        summary_ws["A2"] = "Parameter"
        summary_ws["B2"] = "Single Person"
        summary_ws["C2"] = "Multi Person"

        for cell in summary_ws[2]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        rows = {
            4: "KONFIGURASI PENGUJIAN",
            5: "Waktu Mulai",
            6: "Resolusi Input Aktual",
            7: "Jumlah Gerakan",
            8: "Warm-up (s)",

            10: "LATENSI PEMROSESAN END-TO-END",
            11: "Kode",
            12: "Durasi Workout (s)",
            13: "Jumlah Frame Dianalisis",
            14: "Rata-Rata Latensi (ms)",
            15: "Median Latensi (ms)",
            16: "P95 Latensi (ms)",
            17: "Minimum Latensi (ms)",
            18: "Maksimum Latensi (ms)",

            20: "FRAME RATE EFEKTIF",
            21: "Kode",
            22: "Durasi Workout (s)",
            23: "FPS Input Aktual",
            24: "Jumlah Frame Masuk",
            25: "Jumlah Frame Diproses",
            26: "FPS Efektif",
            27: "Frame Skipping Ratio (%)",
            28: "Frekuensi Inferensi Model (Hz)",

            30: "PENGGUNAAN SUMBER DAYA",
            31: "Kode",
            32: "Durasi Workout (s)",
            33: "CPU Rata-Rata (%)",
            34: "CPU Maksimum (%)",
            35: "GPU Rata-Rata (%)",
            36: "GPU Maksimum (%)",
            37: "RAM Rata-Rata (MB)",
            38: "RAM Maksimum (MB)",
            39: "Memori GPU Maksimum (MB)",
        }

        for row, label in rows.items():
            summary_ws.cell(row=row, column=1, value=label)
        for row in (4, 10, 20, 30):
            summary_ws.cell(row=row, column=1).font = Font(bold=True)

        summary_ws.column_dimensions["A"].width = 40
        summary_ws.column_dimensions["B"].width = 22
        summary_ws.column_dimensions["C"].width = 22

        for mode, payload in self._completed_results.items():
            column = 3 if mode == "multi" else 2
            summary = payload["summary"]
            metadata = payload["metadata"]
            latency_code = "TNF-02" if mode == "multi" else "TNF-01"

            values = {
                5: summary.get("started_at"),
                6: metadata.get("camera_resolution"),
                7: metadata.get("exercise_count"),
                8: metadata.get("warmup_seconds"),

                11: latency_code,
                12: summary.get("duration_s"),
                13: summary.get("processed_frames"),
                14: summary.get("latency_mean_ms"),
                15: summary.get("latency_median_ms"),
                16: summary.get("latency_p95_ms"),
                17: summary.get("latency_min_ms"),
                18: summary.get("latency_max_ms"),

                21: "TNF-03",
                22: summary.get("duration_s"),
                23: summary.get("input_fps"),
                24: summary.get("input_frames"),
                25: summary.get("processed_frames"),
                26: summary.get("effective_fps"),
                27: summary.get("frame_skipping_ratio_percent"),
                28: summary.get("model_inference_hz"),

                31: "TNF-04",
                32: summary.get("duration_s"),
                33: summary.get("cpu_mean_percent"),
                34: summary.get("cpu_max_percent"),
                35: summary.get("gpu_mean_percent"),
                36: summary.get("gpu_max_percent"),
                37: summary.get("ram_mean_mb"),
                38: summary.get("ram_max_mb"),
                39: summary.get("gpu_memory_max_mb"),
            }

            for row, value in values.items():
                summary_ws.cell(
                    row=row,
                    column=column,
                    value=self._round_or_none(value),
                )

            suffix = "Multi" if mode == "multi" else "Single"

            latency_ws = workbook.create_sheet(f"Latency_{suffix}")
            latency_ws.append(["Frame Dianalisis", "Latensi End-to-End (ms)"])
            for index, latency in enumerate(payload["latencies_ms"], start=1):
                latency_ws.append([index, round(float(latency), 4)])

            resource_ws = workbook.create_sheet(f"Resource_{suffix}")
            resource_ws.append(
                [
                    "Waktu Aktif (s)",
                    "CPU Aplikasi (%)",
                    "GPU Utilization Device (%)",
                    "RAM Aplikasi (MB)",
                    "Memori GPU Device (MB)",
                    "Memori GPU PyTorch Reserved (MB)",
                ]
            )
            for sample in payload["resources"]:
                resource_ws.append(
                    [
                        self._round_or_none(sample.get("active_time_s")),
                        self._round_or_none(sample.get("cpu_percent")),
                        self._round_or_none(sample.get("gpu_percent")),
                        self._round_or_none(sample.get("ram_mb")),
                        self._round_or_none(sample.get("gpu_memory_mb")),
                        self._round_or_none(sample.get("torch_gpu_memory_mb")),
                    ]
                )

        temp_path = self.output_path.with_suffix(".tmp.xlsx")
        try:
            workbook.save(temp_path)
            temp_path.replace(self.output_path)
            return True
        except PermissionError as exc:
            print(
                "[PERF_XLSX_SAVE_ERROR]",
                "Close the workbook if it is open in Excel.",
                exc,
            )
            return False
        finally:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass

    def _write_json_fallback(self):
        if self.fallback_json_path is None:
            return None

        serializable = {
            "run_id": self._run_id,
            "sessions": self._completed_results,
        }
        self.fallback_json_path.write_text(
            json.dumps(serializable, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        return self.fallback_json_path

    def finish_and_save(self) -> Path | None:
        if not self.running or self.mode is None:
            return None

        if self.segment_active:
            self._sample_resources()
        self.pause()
        self.resource_timer.stop()

        payload = {
            "metadata": dict(self.metadata),
            "summary": self._build_summary(),
            "latencies_ms": list(self.latency_samples_ms),
            "resources": list(self.resource_samples),
        }
        self._completed_results[self.mode] = payload
        self._save_state()

        saved_path: Path | None = None
        if self._write_xlsx():
            saved_path = self.output_path
        else:
            saved_path = self._write_json_fallback()
            if Workbook is None:
                print(
                    "[PERF_XLSX_DISABLED] openpyxl is not available; "
                    "saved one JSON report instead.",
                )

        print(
            "[PERF_TEST_SAVED]",
            f"mode={self.mode}",
            saved_path,
        )

        self.running = False
        self.segment_active = False
        return saved_path